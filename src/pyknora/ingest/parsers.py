"""文档解析：把上传字节转成纯文本（按页标记分隔）。

支持 txt / md / html / pdf / docx。解析只做「取文本」，不猜结构；失败抛
:class:`ParseError`，由上层把文档状态置为 ``failed`` 并记录原因。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import ClassVar

PAGE_BREAK = "\f"

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".log", ".csv", ".json", ".yaml", ".yml"}
HTML_SUFFIXES = {".html", ".htm"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | HTML_SUFFIXES | {".pdf", ".docx"}

MIME_BY_SUFFIX = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ParseError(Exception):
    """解析失败（格式不支持、内容为空、依赖解析库报错）。"""


@dataclass(slots=True)
class ParsedDocument:
    """解析结果：``text`` 用 ``\\f`` 分页，``pages`` 为页数。"""

    text: str
    pages: int
    mime_type: str
    warnings: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not self.text.strip().strip(PAGE_BREAK)


def suffix_of(filename: str) -> str:
    return Path(filename).suffix.lower()


def is_supported(filename: str) -> bool:
    return suffix_of(filename) in SUPPORTED_SUFFIXES


def guess_mime(filename: str, provided: str | None = None) -> str:
    if provided and provided != "application/octet-stream":
        return provided
    return MIME_BY_SUFFIX.get(suffix_of(filename), "application/octet-stream")


def _decode_text(data: bytes) -> tuple[str, str | None]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(encoding), None if encoding == "utf-8" else f"以 {encoding} 解码"
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "含无法解码字节，已用替换字符"


class _TextExtractor(HTMLParser):
    """把 HTML 压成纯文本：丢脚本样式，块级标签换行。"""

    _BLOCK: ClassVar[set[str]] = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "section",
        "article",
        "header",
        "footer",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "blockquote",
        "pre",
        "table",
    }
    _SKIP: ClassVar[set[str]] = {"script", "style", "noscript", "template"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self._parts.append(data)

    def text(self) -> str:
        return "".join(self._parts)


def _parse_html(data: bytes) -> tuple[str, list[str]]:
    raw, warning = _decode_text(data)
    parser = _TextExtractor()
    parser.feed(raw)
    parser.close()
    return clean_text(parser.text()), [warning] if warning else []


def _parse_text(data: bytes) -> tuple[str, list[str]]:
    raw, warning = _decode_text(data)
    return clean_text(raw), [warning] if warning else []


def _parse_pdf(data: bytes) -> tuple[str, list[str]]:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError as exc:  # pragma: no cover - 依赖已声明
        raise ParseError("缺少 pypdf，无法解析 PDF") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"PDF 解析失败：{exc}") from exc

    warnings: list[str] = []
    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            pages.append(clean_text(page.extract_text() or ""))
        except Exception as exc:
            warnings.append(f"第 {index + 1} 页提取失败：{exc}")
            pages.append("")
    if not any(pages):
        raise ParseError("PDF 未提取到文本（可能是扫描件，需 OCR，P1 范围）")
    return PAGE_BREAK.join(pages), warnings


def _parse_docx(data: bytes) -> tuple[str, list[str]]:
    try:
        from docx import Document as DocxDocument
    except ModuleNotFoundError as exc:  # pragma: no cover - 依赖已声明
        raise ParseError("缺少 python-docx，无法解析 DOCX") from exc
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"DOCX 解析失败：{exc}") from exc

    blocks: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return clean_text("\n".join(blocks)), []


_WS_RE = re.compile(r"[ \t\u3000]+")
_BLANK_LINES_RE = re.compile(r"\n{3,}")


def clean_text(text: str) -> str:
    """统一空白：去行尾空格、压缩连续空行，但保留 ``\\f`` 分页符。"""

    pages = text.split(PAGE_BREAK)
    cleaned = []
    for page in pages:
        page = _WS_RE.sub(" ", page.replace("\r\n", "\n").replace("\r", "\n"))
        page = "\n".join(line.strip() for line in page.split("\n"))
        cleaned.append(_BLANK_LINES_RE.sub("\n\n", page).strip("\n"))
    return PAGE_BREAK.join(cleaned)


def parse_bytes(data: bytes, filename: str, mime_type: str | None = None) -> ParsedDocument:
    """按扩展名解析；不支持的扩展名直接抛 :class:`ParseError`。"""

    suffix = suffix_of(filename)
    if suffix not in SUPPORTED_SUFFIXES:
        supported = "、".join(sorted(SUPPORTED_SUFFIXES))
        raise ParseError(f"不支持的格式 {suffix or '（无扩展名）'}；支持：{supported}")
    if not data:
        raise ParseError("文件内容为空")

    if suffix in TEXT_SUFFIXES:
        text, warnings = _parse_text(data)
    elif suffix in HTML_SUFFIXES:
        text, warnings = _parse_html(data)
    elif suffix == ".pdf":
        text, warnings = _parse_pdf(data)
    else:
        text, warnings = _parse_docx(data)

    result = ParsedDocument(
        text=text,
        pages=text.count(PAGE_BREAK) + 1,
        mime_type=guess_mime(filename, mime_type),
        warnings=warnings,
    )
    if result.is_empty:
        raise ParseError("解析后无有效文本")
    return result
